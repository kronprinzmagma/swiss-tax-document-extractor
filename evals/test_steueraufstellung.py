"""Eval-Test für ``scripts/build_steueraufstellung.py``.

Deterministisch (kein ``llm_full``-Marker), läuft im Default-Lauf
``pytest -m "not llm_full"``. Baut ein synthetisches ``_results_full.json``-
Fixture (NUR Fantasy-Namen Hans/Maria/Lina/Tim — Privacy) und prüft die
erzeugte xlsx via openpyxl zurück:

- Blattstruktur + Spaltenköpfe.
- A/B-Konten-Split (verrechnungssteuer > 0 vs. leer).
- Personen-Pivot (zwei KK-Belege Maria → eine Zeile, KVG-Summe).
- Summenzeilen = Summe der float-Datenzellen.
- Manual-Review: anchor_valid=False / manual_review:* → Status "manuell prüfen",
  Betragszelle None.
- extract_failed-Beleg taucht im Out-of-Scope-Blatt auf (verschwindet nicht).
- Beträge sind float-Zellen.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

import scripts.build_tax_output as bto
import scripts.build_steueraufstellung as bsa
from scripts.build_steueraufstellung import build_workbook, _to_float

_FAMILY_TEST = Path(__file__).parent / "family.yaml.test"


@pytest.fixture(autouse=True)
def _family_maps(monkeypatch):
    """Aktiviert das Rollen-Mapping über die synthetische ``family.yaml.test``.

    Die Vornamen→Rolle-Maps in ``build_tax_output`` werden zur Importzeit aus
    ``ROOT/family.yaml`` gebaut — die existiert im Test-/CI-Lauf nicht. Wir
    bauen sie hier deterministisch aus dem PII-freien ``family.yaml.test``
    (Hans/Maria/Lina/Tim Muster) neu und spiegeln auch den abgeleiteten
    ``_ROLE_TO_CANONICAL`` in ``build_steueraufstellung``.
    """
    from extractors.family import load_family

    fmly = load_family(_FAMILY_TEST)
    assert fmly is not None, "family.yaml.test muss ladbar sein"

    pool = ["Hans", "Maria", "Lina", "Tim", "Anna", "Peter"]
    fantasy: dict[str, str] = {}
    real: dict[str, set[str]] = {}
    from extractors.person_inference import _name_tokens

    for idx, m in enumerate(fmly.members):
        role = bto._FAMILY_ROLE_MAP.get(m.role, bto.ROLE_UNKNOWN)
        if idx < len(pool):
            fantasy[pool[idx]] = role
        for tok in _name_tokens(m.first_name):
            real.setdefault(tok, set()).add(role)

    monkeypatch.setattr(bto, "_FANTASY_FIRSTNAME_TO_ROLE", fantasy)
    monkeypatch.setattr(bto, "_REAL_FIRSTNAME_TO_ROLE", real)
    monkeypatch.setattr(bsa, "_FANTASY_FIRSTNAME_TO_ROLE", fantasy)
    monkeypatch.setattr(bsa, "_REAL_FIRSTNAME_TO_ROLE", real)
    monkeypatch.setattr(bsa, "_ROLE_TO_CANONICAL", bsa._build_role_to_canonical())


def _field(feld: str, value: str, anchor_valid: bool = True) -> dict:
    return {
        "feld": feld,
        "value": value,
        "bbox": [0, 0, 10, 10],
        "page": 1,
        "snippet": f"{feld} {value}",
        "anchor_valid": anchor_valid,
        "inference_source": None,
    }


@pytest.fixture
def entries() -> list[dict]:
    """Synthetisches _results_full.json-Fixture (nur Fantasy-Namen)."""
    return [
        # Bank mit Verrechnungssteuer > 0 → Konten A
        {
            "pdf_name": "bank_hans.pdf", "status": "ok",
            "belegtyp": "bank_zinsausweis",
            "fields": [
                _field("institut", "BANK-A"),
                _field("kontotyp", "Sparkonto"),
                _field("kontoinhaber_name", "Hans Muster"),
                _field("vermoegensstand_3112", "10'000.00"),
                _field("bruttoertrag", "100.00"),
                _field("verrechnungssteuer", "35.00"),
                _field("jahr", "2022"),
            ],
        },
        # Bank ohne Verrechnungssteuer → Konten B
        {
            "pdf_name": "bank_maria.pdf", "status": "ok",
            "belegtyp": "bank_zinsausweis",
            "fields": [
                _field("institut", "BANK-B"),
                _field("kontoinhaber_name", "Maria Muster"),
                _field("vermoegensstand_3112", "5'000.00"),
                _field("bruttoertrag", "0.00"),
                _field("verrechnungssteuer", ""),
                _field("jahr", "2022"),
            ],
        },
        # KK Maria #1
        {
            "pdf_name": "kk_maria_1.pdf", "status": "ok",
            "belegtyp": "kk_praemienbescheinigung",
            "fields": [
                _field("kasse", "KK-B"),
                _field("versicherte_person_name", "Maria Muster"),
                _field("praemie_kvg_total", "1'200.00"),
                _field("jahr", "2022"),
            ],
        },
        # KK Maria #2 → Pivot-Summe 1500.00, EINE Maria-Zeile.
        # Andere Schreibweise ("Muster, Maria") muss DERSELBEN Rolle-Zeile
        # zugeordnet werden (kanonischer Rollen-Schlüssel, nicht roher String).
        {
            "pdf_name": "kk_maria_2.pdf", "status": "ok",
            "belegtyp": "kk_praemienbescheinigung",
            "fields": [
                _field("kasse", "KK-B"),
                _field("versicherte_person_name", "Muster, Maria"),
                _field("praemie_kvg_total", "300.00"),
                _field("jahr", "2022"),
            ],
        },
        # Krankheitskosten Tim
        {
            "pdf_name": "krankheit_tim.pdf", "status": "ok",
            "belegtyp": "krankheitskosten",
            "fields": [
                _field("leistungserbringer", "LEISTUNGSERBRINGER-A"),
                _field("person_name", "Tim Muster"),
                _field("betrag", "250.00"),
                _field("jahr", "2022"),
            ],
        },
        # Krankheitskosten Lina mit fehlendem Anker → manuell prüfen
        {
            "pdf_name": "krankheit_lina.pdf", "status": "ok",
            "belegtyp": "krankheitskosten",
            "fields": [
                _field("leistungserbringer", "LEISTUNGSERBRINGER-A"),
                _field("person_name", "Lina Muster"),
                _field("betrag", "manual_review:no_anchor_betrag", anchor_valid=False),
                _field("jahr", "2022"),
            ],
        },
        # Säule 3a Hans
        {
            "pdf_name": "saeule3a_hans.pdf", "status": "ok",
            "belegtyp": "saeule_3a",
            "fields": [
                _field("stiftung", "STIFTUNG_3A-A"),
                _field("kontoinhaber_name", "Hans Muster"),
                _field("einzahlung_betrag", "7'056.00"),
                _field("jahr", "2022"),
            ],
        },
        # Wertschriften plausibel: Bestand > Bruttoertrag → auto
        {
            "pdf_name": "wsv_ok.pdf", "status": "ok",
            "belegtyp": "wertschriftenverzeichnis",
            "fields": [
                _field("institut", "INST-A"),
                _field("bestand_3112", "50'000.00"),
                _field("bruttoertrag_total", "400.00"),
                _field("jahr", "2022"),
            ],
        },
        # Wertschriften IMPLAUSIBEL: Bestand < Bruttoertrag → manuell prüfen
        {
            "pdf_name": "wsv_implausibel.pdf", "status": "ok",
            "belegtyp": "wertschriftenverzeichnis",
            "fields": [
                _field("institut", "INST-B"),
                _field("bestand_3112", "100.00"),
                _field("bruttoertrag_total", "9'999.00"),
                _field("jahr", "2022"),
            ],
        },
        # Lohnausweis Hans (mit Periode + AHV/ALV/NBU = SPECS-mandatory)
        {
            "pdf_name": "lohn_hans.pdf", "status": "ok",
            "belegtyp": "lohnausweis",
            "fields": [
                _field("arbeitgeber", "ARBEITGEBER-A"),
                _field("arbeitnehmer_name", "Hans Muster"),
                _field("periode_von", "01.01.2022"),
                _field("periode_bis", "31.12.2022"),
                _field("bruttolohn_pos8", "95'400.00"),
                _field("ahv_alv_nbu_abzug_pos9", "5'913.30"),
                _field("nettolohn_pos11", "82'486.70"),
                _field("jahr", "2022"),
            ],
        },
        # Kinderbetreuung mit unbestimmter Person → personensensitives Blatt,
        # ROLE_UNKNOWN darf NIE Status 'auto' bekommen.
        {
            "pdf_name": "kita_unbekannt.pdf", "status": "ok",
            "belegtyp": "kinderbetreuung",
            "fields": [
                _field("anbieter", "KITA-A"),
                _field("kind_name", "manual_review:kind_ambiguous", anchor_valid=False),
                _field("betrag", "3'000.00"),
                _field("jahr", "2022"),
            ],
        },
        # extract_failed-Beleg → Out of Scope, verschwindet nicht
        {
            "pdf_name": "kaputt.pdf", "status": "extract_failed",
            "belegtyp": "lohnausweis",
            "fields": [],
            "error": "JSON-EOF",
        },
    ]


@pytest.fixture
def wb(entries):
    return build_workbook(entries)


def _sheet_dict(ws) -> dict:
    """Worksheet als Liste von Zeilen (list of tuples)."""
    return [tuple(c.value for c in row) for row in ws.iter_rows()]


# ── Unit: _to_float ───────────────────────────────────────────────────────────

def test_to_float_ch_format():
    assert _to_float("1'234.50") == 1234.5
    assert _to_float("10'000.00") == 10000.0
    assert _to_float("250.00") == 250.0


def test_to_float_none_cases():
    assert _to_float(None) is None
    assert _to_float("") is None
    assert _to_float("manual_review:x") is None
    assert _to_float("nicht angegeben") is None
    assert _to_float("kein betrag") is None


def test_to_float_dot_thousands_geparst():
    """Dot-Tausender aus der Extraktion ("7.793" für 7793) wird seit dem
    Punkt-Tausender-Fix in parse_swiss_amount als Tausendertrennung erkannt
    (3 Ziffern nach Punkt — CHF hat max. 2 Rappen-Stellen)."""
    assert _to_float("7.793") == 7793.0
    # Apostroph-Tausender ohne Dezimalstellen bleibt korrekt parsebar.
    assert _to_float("12'211") == 12211.0
    # Echt unparsebare Werte degradieren weiterhin sicher zu None.
    assert _to_float("7.7935") is None


# ── Blattstruktur ─────────────────────────────────────────────────────────────

def test_expected_sheets_present(wb):
    names = set(wb.sheetnames)
    assert "Konten A" in names
    assert "Konten B" in names
    assert "Versicherungsprämien" in names
    assert "Krankheits- und Unfallkosten" in names
    assert "Säule 3a" in names
    assert "Out of Scope" in names


def test_konten_headers(wb):
    ws = wb["Konten A"]
    headers = [c.value for c in ws[2]]
    assert headers == ["Konto", "Abschluss", "Zinsertrag", "Quelle", "Status"]


def test_kk_headers(wb):
    ws = wb["Versicherungsprämien"]
    headers = [c.value for c in ws[2]]
    assert headers == [
        "Person", "Grundversicherung KVG", "Zusatzversicherung VVG",
        "Prämienverbilligung", "Quelle", "Status",
    ]


# ── A/B-Split ─────────────────────────────────────────────────────────────────

def test_ab_split(wb):
    a_rows = _sheet_dict(wb["Konten A"])
    b_rows = _sheet_dict(wb["Konten B"])
    # Konten A enthält den Hans-Beleg (VSt 35 > 0)
    a_quellen = [r[3] for r in a_rows]
    assert "bank_hans.pdf" in a_quellen
    assert "bank_maria.pdf" not in a_quellen
    # Konten B enthält den Maria-Beleg (VSt leer)
    b_quellen = [r[3] for r in b_rows]
    assert "bank_maria.pdf" in b_quellen


def test_konten_amounts_are_float(wb):
    ws = wb["Konten A"]
    # Datenzeile 3: Abschluss (col 2) + Zinsertrag (col 3) sind float
    abschluss = ws.cell(row=3, column=2).value
    zins = ws.cell(row=3, column=3).value
    assert isinstance(abschluss, (int, float))
    assert isinstance(zins, (int, float))
    assert abschluss == 10000.0
    assert zins == 100.0


# ── Personen-Pivot ────────────────────────────────────────────────────────────

def test_kk_pivot_single_maria_row(wb):
    ws = wb["Versicherungsprämien"]
    rows = _sheet_dict(ws)
    # Datenzeilen (ab 3) ohne Total-Zeile zählen die Personen.
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    maria_rows = [r for r in data_rows if r[0] and "Maria" in str(r[0])]
    assert len(maria_rows) == 1, f"erwartet EINE Maria-Zeile, fand {len(maria_rows)}"
    # KVG-Summe = 1200 + 300 = 1500
    kvg = maria_rows[0][1]
    assert isinstance(kvg, (int, float))
    assert kvg == 1500.0


def test_kk_quellen_kommasepariert(wb):
    ws = wb["Versicherungsprämien"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    maria_rows = [r for r in data_rows if "Maria" in str(r[0])]
    quelle = maria_rows[0][4]  # Quelle-Spalte (col 5)
    assert "kk_maria_1.pdf" in quelle
    assert "kk_maria_2.pdf" in quelle
    assert "," in quelle


# ── Summenzeilen ──────────────────────────────────────────────────────────────

def test_kk_total_row(wb):
    ws = wb["Versicherungsprämien"]
    rows = _sheet_dict(ws)
    total_rows = [r for r in rows if r[0] == "Total"]
    assert len(total_rows) == 1
    # KVG-Total = nur Maria (1500), Tim/Lina sind nicht KK
    assert total_rows[0][1] == 1500.0


def test_konten_a_total(wb):
    ws = wb["Konten A"]
    rows = _sheet_dict(ws)
    total_rows = [r for r in rows if r[0] == "Total"]
    assert len(total_rows) == 1
    assert total_rows[0][1] == 10000.0  # Abschluss-Summe
    assert total_rows[0][2] == 100.0    # Zinsertrag-Summe


# ── Wertschriften-Plausibilität ───────────────────────────────────────────────

def test_wertschriften_implausibel_manuell(wb):
    """Bestand < Bruttoertrag im selben Beleg → Status 'manuell prüfen',
    Betragszellen None (kein falscher Summenbeitrag)."""
    ws = wb["Wertschriften"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    implausibel = [r for r in data_rows if "INST-B" in str(r[0])]
    assert len(implausibel) == 1
    assert implausibel[0][-1] == "manuell prüfen"
    # Bestand (col 2) + Bruttoertrag (col 3) None.
    assert implausibel[0][1] is None
    assert implausibel[0][2] is None


def test_wertschriften_plausibel_auto(wb):
    ws = wb["Wertschriften"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    plausibel = [r for r in data_rows if "INST-A" in str(r[0])]
    assert len(plausibel) == 1
    assert plausibel[0][-1] == "auto"
    assert plausibel[0][1] == 50000.0
    assert plausibel[0][2] == 400.0


# ── Rollen-Mapping (Regression) ───────────────────────────────────────────────

def test_frau_maps_to_elternteil_2():
    """Regression Bug (c): die Frau (family.yaml-frau-Vorname) darf NIE auf
    elternteil_1 gemappt werden."""
    role = bto.map_person_to_role("Maria Muster")
    assert role == "elternteil_2"
    assert role != "elternteil_1"
    # Auch mit Anrede-Präfix bleibt es elternteil_2.
    assert bto.map_person_to_role("Frau Maria Muster") == "elternteil_2"


def test_maria_pivot_row_is_elternteil_2(wb):
    """Im KK-Pivot wird Maria als elternteil_2 geführt (nie elternteil_1)."""
    ws = wb["Versicherungsprämien"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    # Keine Zeile darf elternteil_1 für Maria zeigen.
    assert not any("elternteil_1" in str(r[0]) for r in data_rows)


# ── ROLE_UNKNOWN nie 'auto' auf personensensitivem Blatt ──────────────────────

def test_kinderbetreuung_unknown_person_never_auto(wb):
    """Kinderbetreuung (personensensitiv) mit ROLE_UNKNOWN → 'manuell prüfen',
    nie 'auto'; Betragszelle None."""
    ws = wb["Kinderbetreuung"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    unbekannt = [r for r in data_rows if r[0] == bto.ROLE_UNKNOWN]
    assert len(unbekannt) == 1
    assert unbekannt[0][-1] == "manuell prüfen"
    # Betrag (col 3) None — keine Auto-Summe für unbestimmte Person.
    assert unbekannt[0][2] is None


# ── Lohn-Blatt: SPECS-mandatory-Spalten ───────────────────────────────────────

def test_lohn_headers_cover_specs_mandatory(wb):
    """Lohn-Blatt deckt die SPECS-mandatory-Felder ab, die fachlich ins Blatt
    gehören (Periode + AHV/ALV/NBU)."""
    ws = wb["Lohnausweise"]
    headers = [c.value for c in ws[2]]
    joined = " ".join(str(h) for h in headers)
    assert "Periode von" in joined
    assert "Periode bis" in joined
    assert "AHV" in joined  # AHV/ALV/NBU (Pos. 9)


def test_lohn_periode_values(wb):
    ws = wb["Lohnausweise"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    hans = [r for r in data_rows if "Hans" in str(r[0])]
    assert len(hans) == 1
    joined = " ".join(str(c) for c in hans[0] if c is not None)
    assert "01.01.2022" in joined
    assert "31.12.2022" in joined


# ── Manual-Review ─────────────────────────────────────────────────────────────

def test_manual_review_lina(wb):
    ws = wb["Krankheits- und Unfallkosten"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    lina = [r for r in data_rows if "Lina" in str(r[0])]
    assert len(lina) == 1
    # Status (letzte Spalte) "manuell prüfen", Betragszelle (col 2) None
    assert lina[0][-1] == "manuell prüfen"
    assert lina[0][1] is None


def test_tim_auto(wb):
    ws = wb["Krankheits- und Unfallkosten"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    tim = [r for r in data_rows if "Tim" in str(r[0])]
    assert len(tim) == 1
    assert tim[0][1] == 250.0
    assert tim[0][-1] == "auto"


# ── Vollständigkeit: extract_failed verschwindet nicht ────────────────────────

def test_extract_failed_in_oos(wb):
    ws = wb["Out of Scope"]
    rows = _sheet_dict(ws)
    dateien = [r[0] for r in rows]
    assert "kaputt.pdf" in dateien


# ── openpyxl-Roundtrip (valide xlsx) ──────────────────────────────────────────

def test_roundtrip_save_load(wb, tmp_path):
    path = tmp_path / "steueraufstellung.xlsx"
    wb.save(path)
    reloaded = openpyxl.load_workbook(path)
    assert "Konten A" in reloaded.sheetnames
    # Beträge bleiben nach Roundtrip float.
    ws = reloaded["Konten A"]
    assert isinstance(ws.cell(row=3, column=2).value, (int, float))


def test_buffer_roundtrip(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    reloaded = openpyxl.load_workbook(buf)
    assert reloaded.sheetnames  # valide xlsx mit ≥1 Blatt


# ── Steuerjahr-Zeile ──────────────────────────────────────────────────────────

def test_steuerjahr_row(wb):
    ws = wb["Konten A"]
    assert ws.cell(row=1, column=1).value == "Steuerjahr 2022"


# ── check_aufstellung.py: privacy-tauglicher xlsx-Validator ───────────────────

def test_check_aufstellung_good_xlsx_exit_0(wb, tmp_path, capsys):
    """Eine sauber erzeugte xlsx → Exit 0, keine FAIL-Befunde."""
    from scripts.check_aufstellung import main

    path = tmp_path / "steueraufstellung.xlsx"
    wb.save(path)
    rc = main(["--file", str(path)])
    out = capsys.readouterr().out
    assert rc == 0, f"erwartet Exit 0, Output:\n{out}"
    assert "FAIL" not in out


def test_check_aufstellung_string_amount_exit_1(wb, tmp_path, capsys):
    """String in einer Betragszelle → Befund nonnumeric_amount_cell, Exit 1."""
    from scripts.check_aufstellung import main

    # Sabotage: in Konten A (Abschluss, col 2) einen String setzen.
    ws = wb["Konten A"]
    ws.cell(row=3, column=2, value="zehntausend")
    path = tmp_path / "kaputt.xlsx"
    wb.save(path)
    rc = main(["--file", str(path)])
    out = capsys.readouterr().out
    assert rc == 1
    assert "nonnumeric_amount_cell" in out


def test_check_aufstellung_total_mismatch_exit_1(wb, tmp_path, capsys):
    """Verfälschte Total-Summe → Befund total_mismatch, Exit 1."""
    from scripts.check_aufstellung import main

    ws = wb["Konten A"]
    # Total-Zeile finden (Spalte 1 == "Total") und Summenzelle verfälschen.
    for row in ws.iter_rows():
        if row[0].value == "Total":
            ws.cell(row=row[0].row, column=2, value=999999.0)
            break
    path = tmp_path / "total_kaputt.xlsx"
    wb.save(path)
    rc = main(["--file", str(path)])
    out = capsys.readouterr().out
    assert rc == 1
    assert "total_mismatch" in out


def test_check_aufstellung_marker_in_value_exit_1(wb, tmp_path, capsys):
    """manual_review:*-Marker in einer Wertzelle → marker_in_value_cell, Exit 1."""
    from scripts.check_aufstellung import main

    ws = wb["Versicherungsprämien"]
    ws.cell(row=3, column=1, value="manual_review:kaputt")
    path = tmp_path / "marker.xlsx"
    wb.save(path)
    rc = main(["--file", str(path)])
    out = capsys.readouterr().out
    assert rc == 1
    assert "marker_in_value_cell" in out


def test_check_aufstellung_unknown_person_auto_exit_1(wb, tmp_path, capsys):
    """Person 'unbekannt/manuell' mit Status 'auto' → unknown_person_auto, Exit 1."""
    from scripts.check_aufstellung import main

    ws = wb["Versicherungsprämien"]
    # Person-Spalte (col 1) auf ROLE_UNKNOWN, Status-Spalte (letzte) auf auto.
    last_col = ws.max_column
    ws.cell(row=3, column=1, value=bto.ROLE_UNKNOWN)
    ws.cell(row=3, column=last_col, value="auto")
    path = tmp_path / "unknown_auto.xlsx"
    wb.save(path)
    rc = main(["--file", str(path)])
    out = capsys.readouterr().out
    assert rc == 1
    assert "unknown_person_auto" in out


def test_check_aufstellung_output_is_privacy_safe(wb, tmp_path, capsys):
    """Der Output enthält NIE Zellwerte (Beträge/Namen) — nur Befundarten/Zähler."""
    from scripts.check_aufstellung import main

    # Sabotage mit erkennbarem Wert, der NICHT im Output landen darf.
    ws = wb["Konten A"]
    ws.cell(row=3, column=2, value="GEHEIM-4242")
    path = tmp_path / "privacy.xlsx"
    wb.save(path)
    main(["--file", str(path)])
    out = capsys.readouterr().out
    # Weder der sabotierte Wert noch ein Fantasy-Name darf im Output stehen.
    assert "GEHEIM-4242" not in out
    assert "Hans" not in out
    assert "Maria" not in out


# ── Finding 2/4b: KK-Prämien — keine 0.00-Phantomzeile, by-design neutral ─────

def _kk_entry(pdf_name: str, person: str, fields: list[dict]) -> dict:
    return {
        "pdf_name": pdf_name, "status": "ok",
        "belegtyp": "kk_praemienbescheinigung",
        "fields": [_field("versicherte_person_name", person)] + fields,
    }


def _kk_data_rows(wb):
    ws = wb["Versicherungsprämien"]
    rows = _sheet_dict(ws)
    return [r for r in rows[2:] if r[0] and r[0] != "Total"]


def test_kk_ohne_praemienfelder_keine_auto_0_zeile():
    # Beide Prämienfelder fehlen komplett → keine substanzlose auto-0.00-Zeile.
    entry = _kk_entry("kk_leer.pdf", "Maria Muster", [])
    wb = build_workbook([entry])
    data_rows = _kk_data_rows(wb)
    assert data_rows, "Maria-Zeile erwartet (Beleg verschwindet nicht)"
    maria = [r for r in data_rows if "Maria" in str(r[0])][0]
    status = maria[-1]
    assert status == bsa.STATUS_MANUELL
    # Keine auto-Betragszelle mit 0.00.
    assert maria[1] is None  # KVG
    assert maria[2] is None  # VVG


def test_kk_by_design_kvg_plus_echtes_vvg_bleibt_auto():
    # by-design-KVG ist neutral; ein echtes verankertes VVG → Zeile auto mit VVG.
    entry = _kk_entry("kk_vvg.pdf", "Maria Muster", [
        _field("praemie_kvg_total", "not_in_beleg_by_design"),
        _field("praemie_vvg_total", "1'200.00"),
    ])
    wb = build_workbook([entry])
    data_rows = _kk_data_rows(wb)
    maria = [r for r in data_rows if "Maria" in str(r[0])][0]
    assert maria[-1] == bsa.STATUS_AUTO
    assert maria[2] == 1200.0  # VVG-Wert vorhanden, nicht verdeckt


# ── Finding 4a: by-design-kontotyp nie roh in der Konto-Zelle ─────────────────

def test_konten_by_design_kontotyp_nur_institut():
    entry = {
        "pdf_name": "bank_by_design.pdf", "status": "ok",
        "belegtyp": "bank_zinsausweis",
        "fields": [
            _field("institut", "BANK-A"),
            _field("kontotyp", "not_in_beleg_by_design"),
            _field("kontoinhaber_name", "Hans Muster"),
            _field("vermoegensstand_3112", "10'000.00"),
            _field("bruttoertrag", "100.00"),
            _field("verrechnungssteuer", "35.00"),
        ],
    }
    wb = build_workbook([entry])
    # Konto landet in A (VSt > 0).
    ws = wb["Konten A"]
    rows = _sheet_dict(ws)
    data_rows = [r for r in rows[2:] if r[0] and r[0] != "Total"]
    konto_zelle = data_rows[0][0]
    assert "not_in_beleg_by_design" not in str(konto_zelle)
    assert str(konto_zelle) == "BANK-A"
