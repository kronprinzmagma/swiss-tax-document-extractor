#!/usr/bin/env python3
"""Privacy-tauglicher Validator für eine ``steueraufstellung.xlsx``.

Prüft eine vom :mod:`scripts.build_steueraufstellung` erzeugte xlsx maschinell
auf strukturelle Integrität — OHNE jemals Zellwerte (Beträge, Namen, Quellen)
auszugeben. Der Output besteht ausschliesslich aus Befundarten und Zählern
(harte Privacy-Linie, analog ``grade_run.py``): so kann der Check auch im
produktiven Pfad gegen echte Daten laufen, ohne PII zu leaken.

Geprüfte Befundarten:

* ``nonnumeric_amount_cell`` — eine Betragsspalte enthält einen nicht-numerischen,
  nicht-leeren Wert (z.B. ein String-Betrag) → nicht summierbar.
* ``marker_in_value_cell`` — eine Zelle enthält einen ``manual_review:``-Marker
  oder einen Placeholder-Text ("nicht angegeben" etc.) als Wert.
* ``unknown_person_auto`` — eine Datenzeile mit Person ``unbekannt/manuell`` hat
  Status ``auto`` (verletzt die personensensitive Regel).
* ``total_mismatch`` — die Total-Zeile stimmt in einer Summenspalte nicht mit der
  Summe der Datenzellen überein (Toleranz 0.01).
* ``missing_sheet`` — ein erwartetes Blatt fehlt (nur wenn überhaupt eines der
  inhaltlichen Blätter vorhanden ist; ein leerer Lauf ist kein Fehler).

Exit-Code: 0 wenn keine Befunde, sonst 1.

Aufruf::

    PYTHONPATH=. python scripts/check_aufstellung.py --file evals/samples_real_2022/steueraufstellung.xlsx
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.build_tax_output import ROLE_UNKNOWN

STATUS_AUTO = "auto"

# Spaltenköpfe, die eine Betrags-/Summenspalte markieren (Teilstring-Match,
# case-insensitive). Diese Zellen müssen None oder numerisch sein.
_AMOUNT_HEADER_KEYWORDS = (
    "betrag", "bestand", "ertrag", "zins", "abschluss", "bruttolohn",
    "nettolohn", "ahv", "prämie", "praemie", "kvg", "vvg", "einzahlung",
    "kostenbeteiligung", "grundversicherung", "zusatzversicherung",
    "prämienverbilligung", "praemienverbilligung", "kosten",
)

# Placeholder-/Marker-Substrings, die NIE als Wert in einer Zelle stehen dürfen.
_MARKER_SUBSTRINGS = (
    "manual_review:", "nicht angegeben", "placeholder", "coming soon",
    "todo", "fixme", "n/a", "not_in_beleg_by_design",
)


def _is_amount_header(header: object) -> bool:
    if not isinstance(header, str):
        return False
    low = header.lower()
    return any(kw in low for kw in _AMOUNT_HEADER_KEYWORDS)


def _is_data_sheet(ws) -> bool:
    """True für inhaltliche Blätter mit Steuerjahr-Kopf (Z.1) + Spaltenköpfe (Z.2).

    Das Out-of-Scope-Blatt hat seine Köpfe in Zeile 1 (kein Steuerjahr-Header)
    und wird hier bewusst NICHT als Datenblatt geprüft.
    """
    a1 = ws.cell(row=1, column=1).value
    return isinstance(a1, str) and a1.startswith("Steuerjahr")


# Blätter, die — sofern überhaupt ein Datenblatt existiert — strukturell erwartet
# werden. Bewusst klein gehalten (nur die personensensitiven Kern-Blätter), damit
# ein Lauf ohne z.B. Spenden nicht fälschlich als unvollständig gilt.
_EXPECTED_SHEETS_IF_ANY = ("Versicherungsprämien",)


def check_workbook(wb: openpyxl.Workbook) -> dict[str, int]:
    """Validiert ein Workbook und liefert ``{befundart: anzahl}`` (keine Werte)."""
    findings: dict[str, int] = defaultdict(int)

    data_sheets = [ws for ws in wb.worksheets if _is_data_sheet(ws)]

    for ws in wb.worksheets:
        headers = [c.value for c in ws[2]] if ws.max_row >= 2 else []
        amount_cols = {
            idx for idx, h in enumerate(headers, start=1) if _is_amount_header(h)
        }
        # Status-Spalte = letzte Spalte mit Kopf "Status".
        status_col = None
        for idx, h in enumerate(headers, start=1):
            if isinstance(h, str) and h.strip().lower() == "status":
                status_col = idx

        is_data = _is_data_sheet(ws)

        # Datenzeilen ab Zeile 3 (bei Datenblättern), sonst alle Zeilen.
        start_row = 3 if is_data else 1
        sums: dict[int, float] = {c: 0.0 for c in amount_cols}
        total_row_values: dict[int, float] = {}
        total_row_idx = None

        for row in ws.iter_rows(min_row=start_row):
            first = row[0].value if row else None
            is_total = isinstance(first, str) and first.strip() == "Total"
            if is_total:
                total_row_idx = row[0].row
                for cell in row:
                    if cell.column in amount_cols and isinstance(cell.value, (int, float)):
                        total_row_values[cell.column] = float(cell.value)
                continue

            row_has_content = any(c.value not in (None, "") for c in row)
            if not row_has_content:
                continue

            # Marker-/Placeholder-Check über ALLE Zellen.
            for cell in row:
                val = cell.value
                if isinstance(val, str):
                    low = val.lower()
                    if any(sub in low for sub in _MARKER_SUBSTRINGS):
                        findings["marker_in_value_cell"] += 1

            # Betragszellen müssen None oder numerisch sein.
            for cell in row:
                if cell.column in amount_cols:
                    val = cell.value
                    if val is None:
                        continue
                    if isinstance(val, bool) or not isinstance(val, (int, float)):
                        findings["nonnumeric_amount_cell"] += 1
                    else:
                        sums[cell.column] += float(val)

            # unknown_person_auto: Person == ROLE_UNKNOWN UND Status == auto.
            if is_data and status_col is not None:
                person = first
                status_val = ws.cell(row=row[0].row, column=status_col).value
                if (
                    isinstance(person, str)
                    and person.strip() == ROLE_UNKNOWN
                    and isinstance(status_val, str)
                    and status_val.strip() == STATUS_AUTO
                ):
                    findings["unknown_person_auto"] += 1

        # Total-Abgleich: Summenzelle == Summe der Datenzellen (Toleranz 0.01).
        if total_row_idx is not None:
            for col in amount_cols:
                declared = total_row_values.get(col)
                if declared is None:
                    continue
                if abs(declared - sums.get(col, 0.0)) > 0.01:
                    findings["total_mismatch"] += 1

    # Fehlende erwartete Blätter (nur wenn überhaupt ein Datenblatt existiert).
    if data_sheets:
        present = set(wb.sheetnames)
        for name in _EXPECTED_SHEETS_IF_ANY:
            if name not in present:
                findings["missing_sheet"] += 1

    return dict(findings)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Privacy-tauglicher Validator für steueraufstellung.xlsx"
    )
    parser.add_argument(
        "--file", type=Path,
        default=ROOT / "evals" / "samples_real_2022" / "steueraufstellung.xlsx",
        help="Pfad zur zu prüfenden xlsx",
    )
    args = parser.parse_args(argv)

    path = args.file if args.file.is_absolute() else ROOT / args.file
    if not path.exists():
        print(f"FEHLT: {path} — erst build_steueraufstellung laufen lassen.",
              file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(path, data_only=True)
    findings = check_workbook(wb)

    if not findings:
        print("OK: keine Befunde (xlsx privacy-tauglich + strukturell konsistent)")
        return 0

    # Nur Befundart + Zähler — NIE Zellwerte.
    for befund in sorted(findings):
        print(f"FAIL: {befund} x{findings[befund]}")
    return 1


if __name__ == "__main__":
    sys.exit(main())
